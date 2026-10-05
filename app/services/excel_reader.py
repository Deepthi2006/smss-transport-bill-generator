from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl


class ExcelBillReader:
    """
    Reads the bill table from an SMSS Transport Excel workbook.

    The parser dynamically finds:
    - S.NO header row
    - Bill data rows
    - TOTAL row

    It does NOT depend on fixed row numbers.
    """

    REQUIRED_HEADERS = {
        "s.no",
        "from",
        "to",
        "in no",
        "inv.date",
        "item name",
        "veh-no",
        "dispathced qty(kgs)",
        "fright/per ton",
        "total fright",
    }

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"Excel file not found: {self.file_path}"
            )

    @staticmethod
    def clean_text(value: Any) -> str:
        """Convert a cell value into clean text."""
        if value is None:
            return ""

        return " ".join(str(value).split()).strip()

    @staticmethod
    def normalize_header(value: Any) -> str:
        """Normalize Excel header text for comparison."""
        return ExcelBillReader.clean_text(value).lower()

    @staticmethod
    def convert_date(value: Any) -> str:
        """Convert Excel date into DD/MM/YYYY format."""

        if isinstance(value, datetime):
            return value.strftime("%d/%m/%Y")

        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def to_decimal(value: Any) -> Decimal:
        """Safely convert an Excel numeric value to Decimal."""

        if value is None or value == "":
            return Decimal("0")

        return Decimal(str(value))

    def load_workbook(self):
        """Load the Excel workbook."""

        return openpyxl.load_workbook(
            self.file_path,
            data_only=True
        )

    def find_header_row(self, worksheet) -> int:
        """
        Find the row containing the S.NO header.
        """

        for row_number in range(
            1,
            worksheet.max_row + 1
        ):
            for column_number in range(
                1,
                worksheet.max_column + 1
            ):

                value = self.normalize_header(
                    worksheet.cell(
                        row_number,
                        column_number
                    ).value
                )

                if value == "s.no":
                    return row_number

        raise ValueError(
            "Could not find 'S.NO' header in the Excel file."
        )

    def get_headers(
        self,
        worksheet,
        header_row: int
    ) -> dict[str, int]:
        """
        Create:

        normalized header -> column number
        """

        headers = {}

        for column_number in range(
            1,
            worksheet.max_column + 1
        ):

            value = self.normalize_header(
                worksheet.cell(
                    header_row,
                    column_number
                ).value
            )

            if value:
                headers[value] = column_number

        return headers

    def find_total_row(
        self,
        worksheet,
        header_row: int,
        headers: dict[str, int]
    ) -> int:
        """
        Find the TOTAL row.

        We identify it by:
        - S.NO is empty
        - Quantity contains a number
        - Total Freight contains a number
        """

        sno_column = headers["s.no"]

        quantity_column = headers[
            "dispathced qty(kgs)"
        ]

        total_freight_column = headers[
            "total fright"
        ]

        for row_number in range(
            header_row + 1,
            worksheet.max_row + 1
        ):

            sno_value = worksheet.cell(
                row_number,
                sno_column
            ).value

            quantity_value = worksheet.cell(
                row_number,
                quantity_column
            ).value

            total_freight_value = worksheet.cell(
                row_number,
                total_freight_column
            ).value

            if (
                sno_value in (None, "")
                and isinstance(
                    quantity_value,
                    (int, float)
                )
                and isinstance(
                    total_freight_value,
                    (int, float)
                )
            ):
                return row_number

        raise ValueError(
            "Could not find TOTAL row in the Excel file."
        )

    def extract_rows(
        self,
        worksheet,
        header_row: int,
        total_row: int,
        headers: dict[str, int]
    ) -> list[dict[str, Any]]:
        """
        Extract all numbered bill rows.
        """

        rows = []

        sno_column = headers["s.no"]

        for row_number in range(
            header_row + 1,
            total_row
        ):

            sno_value = worksheet.cell(
                row_number,
                sno_column
            ).value

            # Ignore blank/non-numbered rows.
            if not isinstance(
                sno_value,
                (int, float)
            ):
                continue

            row = {
                "sno": int(sno_value),

                "from": self.clean_text(
                    worksheet.cell(
                        row_number,
                        headers["from"]
                    ).value
                ),

                "to": self.clean_text(
                    worksheet.cell(
                        row_number,
                        headers["to"]
                    ).value
                ),

                "invoice_no": self.clean_text(
                    worksheet.cell(
                        row_number,
                        headers["in no"]
                    ).value
                ),

                "invoice_date": self.convert_date(
                    worksheet.cell(
                        row_number,
                        headers["inv.date"]
                    ).value
                ),

                "item_name": self.clean_text(
                    worksheet.cell(
                        row_number,
                        headers["item name"]
                    ).value
                ),

                "vehicle_no": self.clean_text(
                    worksheet.cell(
                        row_number,
                        headers["veh-no"]
                    ).value
                ),

                "dispatched_qty": self.to_decimal(
                    worksheet.cell(
                        row_number,
                        headers["dispathced qty(kgs)"]
                    ).value
                ),

                "freight_per_ton": self.to_decimal(
                    worksheet.cell(
                        row_number,
                        headers["fright/per ton"]
                    ).value
                ),

                "total_freight": self.to_decimal(
                    worksheet.cell(
                        row_number,
                        headers["total fright"]
                    ).value
                ),
            }

            rows.append(row)

        return rows

    def extract_totals(
        self,
        worksheet,
        total_row: int,
        headers: dict[str, int]
    ) -> dict[str, Decimal]:
        """Extract totals from the TOTAL row."""

        quantity = self.to_decimal(
            worksheet.cell(
                total_row,
                headers["dispathced qty(kgs)"]
            ).value
        )

        freight = self.to_decimal(
            worksheet.cell(
                total_row,
                headers["total fright"]
            ).value
        )

        return {
            "quantity": quantity,
            "freight": freight
        }

    def read(self) -> dict[str, Any]:
        """
        Read the complete bill table.
        """

        workbook = self.load_workbook()

        worksheet = workbook.active

        header_row = self.find_header_row(
            worksheet
        )

        headers = self.get_headers(
            worksheet,
            header_row
        )

        missing_headers = (
            self.REQUIRED_HEADERS
            - set(headers.keys())
        )

        if missing_headers:
            raise ValueError(
                "Missing required headers: "
                + ", ".join(
                    sorted(missing_headers)
                )
            )

        total_row = self.find_total_row(
            worksheet,
            header_row,
            headers
        )

        rows = self.extract_rows(
            worksheet,
            header_row,
            total_row,
            headers
        )

        totals = self.extract_totals(
            worksheet,
            total_row,
            headers
        )

        return {
            "file_name": self.file_path.name,
            "sheet_name": worksheet.title,
            "header_row": header_row,
            "total_row": total_row,
            "rows": rows,
            "totals": totals
        }


def read_excel(
    file_path: str | Path
) -> dict[str, Any]:
    """
    Convenience function for the rest of the application.
    """

    reader = ExcelBillReader(file_path)

    return reader.read()